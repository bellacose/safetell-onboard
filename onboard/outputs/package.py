"""The hand-off package for SafeTell. SafeTell's own importers do the final validation;
this app only prepares clean files in their formats:
  SafeTell-Onboarding-Workbook_<customer>.xlsx  -> Onboarding page (import_onboarding_workbook)
  history.json                                   -> Import history (import_history)
  checklist_<name>.xlsx                          -> Inspection templates > Import Excel
  requirements.json                              -> reviewed requirements (see docs/handoff-format.md)
"""
from __future__ import annotations
import copy, difflib, io, json, os, re, zipfile
import openpyxl

TEMPLATE = os.path.join(os.path.dirname(__file__), 'templates', 'SafeTell-Onboarding-Workbook.xlsx')
SUFFIX = re.compile(r'\b(llc|inc|co|company|corp|corporation|ltd|lp|llp|pllc|group)\b')


def norm_company(s: str) -> str:
    s = re.sub(r'[^a-z0-9 ]+', ' ', (s or '').lower())
    return re.sub(r'\s+', ' ', SUFFIX.sub(' ', s)).strip()


def possible_duplicates(contractors: list[dict]) -> list[dict]:
    """Same firm entered twice: normalized name, close spelling, or same company email domain."""
    free = {'gmail.com', 'yahoo.com', 'hotmail.com', 'outlook.com', 'aol.com', 'icloud.com', 'live.com', 'msn.com', 'comcast.net'}
    out = []
    for i, a in enumerate(contractors):
        for b in contractors[i + 1:]:
            na, nb = norm_company(a.get('company')), norm_company(b.get('company'))
            da = (a.get('contact_email') or '').lower().partition('@')[2]
            db = (b.get('contact_email') or '').lower().partition('@')[2]
            why = ('same name' if na == nb else 'similar name' if difflib.SequenceMatcher(None, na, nb).ratio() >= 0.8
                   else 'same email domain' if da and da == db and da not in free else None)
            if why:
                out.append({'rows': [a['_row'], b['_row']], 'companies': [a.get('company'), b.get('company')], 'why': why})
    return out


def _write_rows(ws, rows: list[list], start_row: int = 3):
    for r, values in enumerate(rows, start=start_row):
        for c, v in enumerate(values, start=1):
            ws.cell(row=r, column=c, value=v)


def build_workbook(*, company: dict | None = None, users: list[dict] | None = None, contractors: list[dict] | None = None) -> bytes:
    """Pre-fill what the customer already gave us; the decisions tabs stay for the kickoff call."""
    wb = openpyxl.load_workbook(TEMPLATE)
    if company:
        ws = wb['Company']
        for row in ws.iter_rows(min_row=2):
            field = (row[0].value or '').strip()
            if field in company and company[field] not in (None, ''):
                row[1].value = company[field]
    if users:
        _write_rows(wb['Users'], [[u.get('full_name'), u.get('email'), u.get('job_title'), u.get('role', ''), u.get('signs_as', ''), u.get('escalations', '')] for u in users])
    if contractors:
        _write_rows(wb['Contractors'], [[c.get('company'), c.get('trade'), c.get('risk_level', ''), c.get('contact_name'), c.get('contact_email'),
                                         c.get('phone'), c.get('city'), c.get('state', ''), 'Y' if c.get('contact_email') else 'N'] for c in contractors])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_history(corrective_actions=None, injuries=None, hours=None, permits=None, past_audits=None) -> dict:
    """Payload for SafeTell's import_history(project, payload, dry_run)."""
    clean = lambda rows: [{k: v for k, v in r.items() if v not in (None, '')} for r in (rows or [])]
    return {'corrective_actions': clean(corrective_actions), 'injuries': clean(injuries), 'hours': clean(hours),
            'permits': clean(permits), 'past_audits': clean(past_audits)}


# Column headers exactly as SafeTell's Import history screen labels them (src/lib/historyImport.ts),
# so its column matching picks every column without the customer touching anything.
HISTORY_LABELS = {
    'corrective_actions': [('ref', 'Ref / number'), ('found_date', 'Date found'), ('contractor', 'Contractor'), ('package_code', 'Package code'),
                           ('location', 'Location'), ('description', 'Description'), ('class', 'Class'), ('due_date', 'Due date'),
                           ('status', 'Status (open/closed)'), ('closed_date', 'Closed date')],
    'injuries': [('case_ref', 'Case number'), ('date', 'Date'), ('employer', 'Employer'), ('recorded_by', 'Recorded by (if different)'),
                 ('employee', 'Employee'), ('job_title', 'Job title'), ('description', 'Description'), ('classification', 'Classification'),
                 ('days_away', 'Days away'), ('days_restricted', 'Days restricted')],
    'hours': [('contractor', 'Contractor'), ('month', 'Month'), ('hours', 'Hours'), ('reported_by', 'Reported by (optional)')],
    'permits': [('number', 'Permit number'), ('type', 'Type'), ('contractor', 'Contractor'), ('location', 'Location / zone'), ('valid_from', 'Valid from'),
                ('valid_to', 'Valid to'), ('status', 'Status'), ('signers', 'Signers'), ('details', 'Details (e.g. ICRA class)')],
}


def build_history_sheet(section: str, rows: list[dict]) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = section
    ws.append([label for _, label in HISTORY_LABELS[section]])
    for r in rows:
        ws.append([r.get(k, '') for k, _ in HISTORY_LABELS[section]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


CHECKLIST_HEADERS = ['Category', 'Category Weight', 'Question', 'Severity', 'Points', 'Requirement Code', 'Life Critical']


def build_checklist(items: list[dict], life_critical_confirmed: set[int]) -> bytes:
    """SafeTell's checklist template format (src/lib/inspectionExcel.ts). Items: {section, text, severity?}."""
    sections = []
    for it in items:
        if it.get('section') not in sections:
            sections.append(it.get('section') or 'General')
    weight = round(100 / len(sections)) if sections else 0
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Checklist'
    ws.append(CHECKLIST_HEADERS)
    for i, it in enumerate(items):
        lc = i in life_critical_confirmed
        ws.append([it.get('section') or 'General', weight, it['text'], 'critical' if lc else (it.get('severity') or 'medium'), '', it.get('requirement_code', ''), 'Yes' if lc else ''])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_zip(files: dict[str, bytes | str]) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data if isinstance(data, (bytes, str)) else json.dumps(data, indent=1))
    return buf.getvalue()
