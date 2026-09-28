"""The Brightwater mock customer is the answer key: these tests fail if the workbench
stops producing what a careful person would."""
import json, os
import pytest
from onboard.models import Engagement, Upload, make_session
from onboard.docs import review
from onboard.docs.extract import extract_text
from onboard.docs.paragraphs import split_paragraphs
from onboard.sheets.tables import read_sheets, detect_header, auto_map, records, missing_required
from onboard.sheets.lifecritical import suggestions
from onboard.outputs import package
from onboard.assessment.forms import score, documents_for_rank
import datetime as dt

F = os.path.join(os.path.dirname(__file__), 'fixtures')
fx = lambda n: os.path.join(F, n)
TRUTH = json.load(open(fx('truth.json')))
LIFE_CRITICAL = ["Fall protection in place at 6 ft", "Excavation protective system at 5 ft", "Energized work permit", "LOTO verified",
                 "Crane lift plan approved and posted", "ICRA barrier intact and negative pressure verified", "Fire watch in place during hot work"]


@pytest.fixture
def s():
    return make_session('sqlite:///:memory:')()


def _doc(s, kind, filename, title):
    e = Engagement(name='Brightwater Builders')
    u = Upload(engagement=e, kind=kind, filename=filename, path=fx(filename))
    s.add_all([e, u]); s.flush()
    return review.create_document(s, e, u, open(fx(filename), 'rb').read(), title), u


# ---------------------------------------------------------------- splitting
def test_manual_splits_into_every_numbered_paragraph():
    ps = split_paragraphs(extract_text('m.docx', open(fx('BWB-Safety-Program-Manual-Rev7.docx'), 'rb').read()))
    refs = {p['ref'] for p in ps if not p.get('heading')}
    assert set(TRUTH['manual_requirements']) <= refs
    assert '3.0' not in refs  # table fragment "3.0 or less" is not a paragraph


def test_spec_keeps_csi_section_numbers():
    ps = split_paragraphs(extract_text('s.pdf', open(fx('Bayside-Health-Spec-01-35-23-and-01-35-13.13.pdf'), 'rb').read()))
    refs = [p['ref'] for p in ps]
    assert 'Section 01 35 23' in refs and 'Section 01 35 13.13' in refs
    assert '01 35 13.13 1.3.A' in refs and '01 35 23 1.3.A' in refs
    assert not any(r.endswith(' (2)') for r in refs)


# ---------------------------------------------------------------- review
def test_review_with_safetell_prompt_output_covers_everything(s):
    doc, u = _doc(s, 'company_manual', 'BWB-Safety-Program-Manual-Rev7.docx', 'BWB Safety Program Manual Rev 7')
    r = review.run_ai(s, doc, u, provider='replay', replay_path=fx('ai_manual.json'))
    assert r['added'] == 61 and not r['rejected']
    c = review.coverage(doc)
    assert c['unverified_quotes'] == 0 and c['number_flags'] == 0
    assert set(k for k, v in c['states'].items() if v == 'not_covered') == set(TRUTH['manual_not_requirements'])


def test_blind_run_misses_are_caught_and_block_completion(s, tmp_path):
    doc, u = _doc(s, 'company_manual', 'BWB-Safety-Program-Manual-Rev7.docx', 'BWB Manual')
    run1 = [{'paragraph_ref': d['citation'], 'title': d['title'], 'requirement_text': d['requirement_text'], 'category': d['category'],
             'high_risk': d['high_risk'], 'quote': d['verbatim_quote']} for d in json.load(open(fx('ai_blind_run1.json'))) if d['source'] == 'manual']
    p = tmp_path / 'run1.json'; p.write_text(json.dumps({'drafts': run1}))
    review.run_ai(s, doc, u, provider='replay', replay_path=str(p))
    c = review.coverage(doc)
    not_cov = {k for k, v in c['states'].items() if v == 'not_covered'}
    assert set(TRUTH['blind_misses']) <= not_cov
    assert c['unverified_quotes'] == 19            # the paraphrased "verbatim" quotes, caught
    with pytest.raises(review.ReviewError, match='not covered'):
        review.complete_review(doc)
    # an unverified quote can't be accepted silently
    bad = next(d for d in doc.drafts if not d.quote_verified)
    with pytest.raises(review.ReviewError, match='Quote not found'):
        review.decide(s, bad, 'accept')
    review.decide(s, bad, 'accept', without_quote=True)
    assert bad.accepted_without_quote


def test_no_ai_documents_are_refused(s):
    doc, u = _doc(s, 'company_manual', 'BWB-Safety-Program-Manual-Rev7.docx', 'BWB Manual')
    u.no_ai = True
    with pytest.raises(review.ReviewError, match='not sent to AI'):
        review.run_ai(s, doc, u, provider='replay', replay_path=fx('ai_manual.json'))


def test_full_review_exports_requirements(s):
    doc, u = _doc(s, 'owner_spec', 'Bayside-Health-Spec-01-35-23-and-01-35-13.13.pdf', 'Bayside Health Spec')
    doc.project_code = 'SAMC-EPT'
    review.run_ai(s, doc, u, provider='replay', replay_path=fx('ai_spec.json'))
    for ref, state in review.coverage(doc)['states'].items():
        if state == 'not_covered':
            review.skip_paragraph(s, doc, ref, 'Boilerplate')
    for d in doc.drafts:
        review.decide(s, d, 'accept')
    review.complete_review(doc)
    out = review.export_requirements(doc)
    assert len(out) == 14 and all(r['layer'] == 'project' and r['project_code'] == 'SAMC-EPT' for r in out)
    assert any(r['source_citation'] == 'Bayside Health Spec 01 35 13.13 1.3.A' for r in out)


# ---------------------------------------------------------------- spreadsheets
@pytest.mark.parametrize('file,sheet,target,n,header', [
    ('BWB_Vendor_List_export_2026-09-22.xlsx', 'Vendors', 'contractors', 10, 3),
    ('SAMC_Corrective_Action_Log_2026-09-26.xlsx', 'CA Log', 'corrective_actions', 14, 2),
    ('SAMC_Injury_Log_and_Manhours_2026.xlsx', 'Injury Log', 'injuries', 4, 2),
    ('SAMC_Injury_Log_and_Manhours_2026.xlsx', 'Manhours', 'hours', 10, 1),
    ('BWB_SF-12_Weekly_Site_Inspection.xlsx', 'SF-12', 'checklist', 40, 4)])
def test_their_spreadsheets_map_without_help(file, sheet, target, n, header):
    sheets = read_sheets(file, open(fx(file), 'rb').read())
    grid = sheets.get(sheet) or next(iter(sheets.values()))
    h = detect_header(grid, target)
    assert h + 1 == header
    m = auto_map(grid[h], target)
    assert not missing_required(m, target)
    assert len(records(grid, h, m)) == n


def test_life_critical_suggestions():
    sheets = read_sheets('c.xlsx', open(fx('BWB_SF-12_Weekly_Site_Inspection.xlsx'), 'rb').read())
    grid = next(iter(sheets.values())); h = detect_header(grid, 'checklist')
    items = records(grid, h, auto_map(grid[h], 'checklist'))
    sug = suggestions(items, LIFE_CRITICAL)
    matched = {x['text'] for x in sug if x['list_match']}
    assert 'Fall protection in use at 6 ft and above' in matched and 'Negative pressure verified (manometer reading logged)' in matched
    assert {x['text'] for x in sug if x['disagree']} == {'Floor holes covered, secured and marked', 'Hot work permit posted'}


def test_vendor_duplicates_flagged():
    sheets = read_sheets('v.xlsx', open(fx('BWB_Vendor_List_export_2026-09-22.xlsx'), 'rb').read())
    grid = sheets['Vendors']; h = detect_header(grid, 'contractors')
    dups = package.possible_duplicates(records(grid, h, auto_map(grid[h], 'contractors')))
    assert any({'Mangrove Electric Co.', 'Mangrove Electrical Company'} == set(d['companies']) for d in dups)


def test_workbook_prefill_reads_back():
    import openpyxl, io
    data = package.build_workbook(company={'Legal name': 'Brightwater Builders Group, Inc.'},
                                  contractors=[{'company': 'Suncoast Concrete Forming LLC', 'contact_email': 'a@b.example'}])
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert wb['Contractors']['A3'].value == 'Suncoast Concrete Forming LLC' and wb['Contractors']['I3'].value == 'Y'
    assert wb['Company']['B2'].value == 'Brightwater Builders Group, Inc.'


# ---------------------------------------------------------------- assessment
def test_brightwater_scores_r3_with_priority_flags():
    a = json.load(open(fx('brightwater_assessment.json')))
    r = score(a, today=dt.date(2026, 9, 28), deadline=dt.date.fromisoformat(a['A5']['deadline_date']))
    assert (r['b1'], r['b2'], r['b3'], r['b4'], r['composite'], r['rank']) == (25, 17, 5, 16, 63, 'R3')
    assert r['required_tier'] == 4 and r['evidence_set'] == 'extended'
    assert len(r['auto_p1_gaps']) == 3
    assert any('deadline' in f for f in r['priority_flags'])
    assert len(documents_for_rank('R3')) == 23


def test_ai_not_run_twice_and_bulk_accept(s):
    doc, u = _doc(s, 'company_manual', 'BWB-Safety-Program-Manual-Rev7.docx', 'BWB Manual')
    review.run_ai(s, doc, u, provider='replay', replay_path=fx('ai_manual.json'))
    with pytest.raises(review.ReviewError, match='already exist'):
        review.run_ai(s, doc, u, provider='replay', replay_path=fx('ai_manual.json'))
    assert review.accept_all_verified(s, doc) == 61


def test_draft_on_unknown_paragraph_is_dropped(s, tmp_path):
    doc, u = _doc(s, 'company_manual', 'BWB-Safety-Program-Manual-Rev7.docx', 'BWB Manual')
    p = tmp_path / 'x.json'; p.write_text(json.dumps({'drafts': [{'paragraph_ref': '99.9', 'title': 't', 'requirement_text': 'r', 'quote': 'q'}]}))
    r = review.run_ai(s, doc, u, provider='replay', replay_path=str(p))
    assert r['added'] == 0 and r['rejected'] and not doc.drafts
