"""Click through Brightwater in the web app (Flask test client), AI replayed from the fixture."""
import io, json, os, zipfile
import openpyxl, pytest
from onboard.web.app import create_app
from onboard.docs import review

F = os.path.join(os.path.dirname(__file__), 'fixtures')


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.delenv('ONBOARD_PASSWORD', raising=False)
    monkeypatch.setattr(review, 'draft_requirements', lambda paras, kind, title, provider=None, replay_path=None:
                        json.load(open(os.path.join(F, 'ai_manual.json' if kind == 'company_manual' else 'ai_spec.json')))['drafts'])
    app = create_app(f"sqlite:///{tmp_path/'t.db'}", str(tmp_path))
    return app.test_client()


def up(c, name, kind, **extra):
    data = {'file': (open(os.path.join(F, name), 'rb'), name), 'kind': kind, **extra}
    return c.post('/e/1/upload', data=data, content_type='multipart/form-data')


def test_brightwater_end_to_end(client):
    c = client
    assert c.post('/engagements', data={'name': 'Brightwater Builders', 'contact_name': 'Marisol Vance', 'contact_email': 'mvance@brightwaterbuilders.example'}).status_code == 302
    # assessment
    ans = open(os.path.join(F, 'brightwater_assessment.json'), 'rb').read()
    c.post('/e/1/assessment/import', data={'file': (io.BytesIO(ans), 'a.json')}, content_type='multipart/form-data')
    page = c.get('/e/1/assessment').get_data(as_text=True)
    assert '63 · R3' in page and 'Priority routing' in page
    # manual: AI, then the review cannot close until skips are recorded
    r = up(c, 'BWB-Safety-Program-Manual-Rev7.docx', 'company_manual', title='BWB Safety Program Manual Rev 7')
    assert '/doc/1' in r.headers['Location']
    c.post('/e/1/doc/1/ai')
    c.post('/e/1/doc/1/complete')
    assert 'not covered' in c.get('/e/1/doc/1').get_data(as_text=True)
    for ref in ('1.1', '18.1'):
        c.post('/e/1/doc/1/skip', data={'ref': ref, 'reason': 'Policy statement'})
    from onboard.models import Draft
    for i in range(1, 62):
        c.post('/e/1/doc/1/decide', data={'draft_id': i, 'decision': 'accept', 'without_quote': '1'})
    c.post('/e/1/doc/1/complete')
    assert 'Review complete' in c.get('/e/1/doc/1').get_data(as_text=True)
    # spreadsheets: default mapping saved as-is
    for name, kind in [('BWB_Vendor_List_export_2026-09-22.xlsx', 'vendor_list'), ('SAMC_Corrective_Action_Log_2026-09-26.xlsx', 'ca_log'), ('BWB_SF-12_Weekly_Site_Inspection.xlsx', 'checklist')]:
        loc = up(c, name, kind).headers['Location']
        page = c.get(loc).get_data(as_text=True)
        assert 'Required and not mapped' not in page
        form = {'save': '1'}
        import re
        for m in re.finditer(r'<select name="(col_\w+)">(.*?)</select>', page, re.S):
            sel = re.search(r'<option selected>(.*?)</option>', m.group(2))
            form[m.group(1)] = sel.group(1) if sel else ''
        form['sheet'] = re.search(r'name="sheet" value="([^"]+)"', page).group(1)
        form['target'] = re.search(r'name="target" value="([^"]+)"', page).group(1)
        form['header'] = re.search(r'name="header" value="(\d+)"', page).group(1)
        form['lc'] = re.findall(r'name="lc" value="(\d+)" checked', page)
        form['exclude'] = re.findall(r'name="exclude" value="(\d+)"', page)
        assert c.post(loc, data=form).status_code == 302
    # package
    z = zipfile.ZipFile(io.BytesIO(c.get('/e/1/package.zip').data))
    names = z.namelist()
    assert any(n.startswith('SafeTell-Onboarding-Workbook_') for n in names)
    assert 'history_corrective_actions.xlsx' in names and 'requirements.json' in names
    assert len(json.loads(z.read('requirements.json'))) == 61
    wb = openpyxl.load_workbook(io.BytesIO(z.read(next(n for n in names if n.startswith('SafeTell-Onboarding')))))
    companies = [r[0].value for r in wb['Contractors'].iter_rows(min_row=3) if r[0].value]
    assert 'Mangrove Electrical Company' not in companies and 'Suncoast Concrete Forming LLC' in companies
    ck = openpyxl.load_workbook(io.BytesIO(z.read(next(n for n in names if n.startswith('checklist_')))))
    lc = [r[2].value for r in ck.active.iter_rows(min_row=2) if r[6].value == 'Yes']
    assert len(lc) >= 9
